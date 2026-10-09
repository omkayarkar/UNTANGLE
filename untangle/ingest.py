"""PDF, vector drawing, OCR, and Excel ingestion with evidence provenance."""

from __future__ import annotations

import io
import math
import re
import shutil
from pathlib import Path
from typing import Any

import pymupdf as fitz
import pandas as pd
from openpyxl import load_workbook

try:
    from .untangle_core import (
        DocumentAnalysis,
        Equipment,
        LineSegment,
        TagOccurrence,
        find_tag_strings,
        normalize_tag,
    )
except ImportError:  # Streamlit also supports running app.py directly.
    from untangle_core import (
        DocumentAnalysis,
        Equipment,
        LineSegment,
        TagOccurrence,
        find_tag_strings,
        normalize_tag,
    )


class IngestionError(ValueError):
    """Input cannot be processed safely or is missing required information."""


EQUIPMENT_WORDS = (
    ("heat exchanger", "Heat exchanger"),
    ("exchanger", "Heat exchanger"),
    ("control valve", "Control valve"),
    ("valve", "Valve"),
    ("pump", "Pump"),
    ("tank", "Tank"),
    ("vessel", "Vessel"),
    ("separator", "Separator"),
    ("filter", "Filter"),
    ("compressor", "Compressor"),
    ("reactor", "Reactor"),
    ("column", "Column"),
)
PREFIX_EQUIPMENT = {
    "P": "Pump",
    "PU": "Pump",
    "PMP": "Pump",
    "TK": "Tank",
    "TANK": "Tank",
    "V": "Vessel",
    "VSL": "Vessel",
    "SEP": "Separator",
    "E": "Heat exchanger",
    "HX": "Heat exchanger",
    "HEX": "Heat exchanger",
    "F": "Filter",
    "C": "Compressor",
    "R": "Reactor",
}


def excel_sheet_names(data: bytes) -> list[str]:
    if not data:
        raise IngestionError("The Excel file is empty.")
    try:
        workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        names = workbook.sheetnames
        workbook.close()
        if not names:
            raise IngestionError("The Excel workbook has no worksheets.")
        return names
    except IngestionError:
        raise
    except Exception as exc:
        raise IngestionError(f"Could not read the Excel workbook: {exc}") from exc


def read_excel_preview(data: bytes, sheet_name: str | int = 0) -> pd.DataFrame:
    excel_sheet_names(data)
    try:
        return pd.read_excel(io.BytesIO(data), sheet_name=sheet_name, engine="openpyxl")
    except Exception as exc:
        raise IngestionError(f"Could not read the selected worksheet: {exc}") from exc


def suggest_tag_column(columns: list[object]) -> str | None:
    """Prefer explicit I/O tag headers; return None rather than guessing."""
    normalized = {str(column).strip().lower(): str(column) for column in columns}
    preferred = (
        "tag",
        "tag no",
        "tag number",
        "instrument tag",
        "instrument tag no",
        "io tag",
        "i/o tag",
        "instrument",
    )
    for name in preferred:
        if name in normalized:
            return normalized[name]
    candidates = [
        original for lowered, original in normalized.items()
        if "tag" in lowered and ("instrument" in lowered or "io" in lowered or "i/o" in lowered)
    ]
    return candidates[0] if len(candidates) == 1 else None


def read_io_values(data: bytes, sheet_name: str, column: str) -> list[object]:
    frame = read_excel_preview(data, sheet_name)
    if column not in frame.columns:
        raise IngestionError(f"Column {column!r} is not present in the selected worksheet.")
    values = frame[column].tolist()
    if not values:
        raise IngestionError("The selected I/O-list column contains no rows.")
    if all(pd.isna(value) or not str(value).strip() for value in values):
        raise IngestionError("The selected I/O-list column contains no tag values.")
    return values


def _span_records(page: fitz.Page) -> list[dict[str, Any]]:
    records: list[dict[str, Any]] = []
    try:
        blocks = page.get_text("dict").get("blocks", [])
    except Exception:
        return records
    for block in blocks:
        for line in block.get("lines", []):
            for span in line.get("spans", []):
                text = str(span.get("text", "")).strip()
                if text:
                    records.append(
                        {
                            "text": text,
                            "bbox": tuple(float(v) for v in span["bbox"]),
                        }
                    )
    return records


def _context_for_bbox(
    bbox: tuple[float, float, float, float], spans: list[dict[str, Any]]
) -> str:
    x = (bbox[0] + bbox[2]) / 2
    y = (bbox[1] + bbox[3]) / 2
    nearby = []
    for span in spans:
        sb = span["bbox"]
        sx = (sb[0] + sb[2]) / 2
        sy = (sb[1] + sb[3]) / 2
        if abs(sy - y) <= 30 and abs(sx - x) <= 130:
            nearby.append((abs(sy - y) * 3 + abs(sx - x), span["text"]))
    nearby.sort(key=lambda item: item[0])
    return " ".join(text for _, text in nearby[:7])


def _rectangles(page: fitz.Page, page_number: int) -> list[tuple[float, float, float, float]]:
    found = []
    try:
        for drawing in page.get_drawings():
            for item in drawing.get("items", []):
                if item[0] == "re":
                    rect = item[1]
                    found.append((float(rect.x0), float(rect.y0), float(rect.x1), float(rect.y1)))
            if any(item[0] == "c" for item in drawing.get("items", [])):
                rect = drawing.get("rect")
                if rect is not None:
                    width, height = float(rect.width), float(rect.height)
                    if 10 <= width <= 260 and 10 <= height <= 180:
                        found.append((float(rect.x0), float(rect.y0), float(rect.x1), float(rect.y1)))
    except Exception:
        return []
    return found


def _symbol_bounds(
    bbox: tuple[float, float, float, float],
    rects: list[tuple[float, float, float, float]],
) -> tuple[float, float, float, float] | None:
    cx, cy = (bbox[0] + bbox[2]) / 2, (bbox[1] + bbox[3]) / 2
    candidates = []
    for rect in rects:
        x0, y0, x1, y1 = rect
        width, height = x1 - x0, y1 - y0
        if not (10 <= width <= 260 and 10 <= height <= 180):
            continue
        dx = max(x0 - cx, 0.0, cx - x1)
        dy = max(y0 - cy, 0.0, cy - y1)
        distance = math.hypot(dx, dy)
        if distance <= 18:
            # Prefer a symbol that contains the text, otherwise the closest
            # nearby shape is only a geometric hint.
            score = (0 if distance < 1 else 1000 + distance * 20) + width * height
            candidates.append((score, rect))
    return min(candidates, default=(0, None), key=lambda item: item[0])[1]


def _equipment_kind(tag: str, context: str) -> tuple[str | None, str, str]:
    lowered = context.lower()
    for word, kind in EQUIPMENT_WORDS:
        if word in lowered:
            return kind, f"Explicit nearby label: {word}", "label-confirmed"
    prefix = re.match(r"[A-Z]+", normalize_tag(tag))
    if prefix:
        kind = PREFIX_EQUIPMENT.get(prefix.group(0))
        if kind:
            return kind, f"Tag-prefix candidate: {prefix.group(0)}", "candidate"
    return None, "No equipment evidence", "unclassified"


def _extract_page_vectors(
    page: fitz.Page, page_number: int, first_path_id: int
) -> tuple[list[LineSegment], int]:
    segments: list[LineSegment] = []
    try:
        drawings = page.get_drawings()
    except Exception:
        return segments, first_path_id
    path_id = first_path_id
    for drawing in drawings:
        for item in drawing.get("items", []):
            if item[0] != "l":
                continue
            start, end = item[1], item[2]
            segment = LineSegment(
                page=page_number,
                x1=float(start.x),
                y1=float(start.y),
                x2=float(end.x),
                y2=float(end.y),
                path_id=path_id,
            )
            if segment.length >= 5:
                segments.append(segment)
        path_id += 1
    return segments, path_id


def _ocr_page(page: fitz.Page, page_number: int) -> list[TagOccurrence]:
    if not shutil.which("tesseract"):
        return []
    try:
        from PIL import Image
        import cv2
        import numpy as np
        import pytesseract
        from pytesseract import Output

        pixmap = page.get_pixmap(matrix=fitz.Matrix(2.2, 2.2), alpha=False)
        pillow_image = Image.frombytes(
            "RGB", (pixmap.width, pixmap.height), pixmap.samples
        )
        image = np.asarray(pillow_image)
        gray = cv2.cvtColor(image, cv2.COLOR_RGB2GRAY)
        cleaned = cv2.threshold(gray, 0, 255, cv2.THRESH_BINARY + cv2.THRESH_OTSU)[1]
        data = pytesseract.image_to_data(
            cleaned, output_type=Output.DICT, config="--psm 11"
        )
    except Exception:
        return []

    sx = page.rect.width / max(1, pixmap.width)
    sy = page.rect.height / max(1, pixmap.height)
    found: list[TagOccurrence] = []
    for index, word in enumerate(data.get("text", [])):
        text = str(word).strip()
        if not text:
            continue
        confidence = -1.0
        try:
            confidence = float(data["conf"][index])
        except (TypeError, ValueError, KeyError):
            pass
        if confidence < 0:
            continue
        for tag in find_tag_strings(text):
            left = int(data["left"][index])
            top = int(data["top"][index])
            width = int(data["width"][index])
            height = int(data["height"][index])
            bbox = (left * sx, top * sy, (left + width) * sx, (top + height) * sy)
            found.append(
                TagOccurrence(
                    original=tag,
                    normalized=normalize_tag(tag),
                    page=page_number,
                    bbox=bbox,
                    source="OCR",
                    confidence=confidence,
                    context=text,
                )
            )
    return found


def inspect_pdf(data: bytes, file_name: str = "drawing.pdf", use_ocr: bool = True) -> DocumentAnalysis:
    if not data:
        raise IngestionError("The PDF is empty.")
    if not data.startswith(b"%PDF"):
        raise IngestionError("The uploaded drawing is not a valid PDF file.")
    try:
        document = fitz.open(stream=data, filetype="pdf")
    except Exception as exc:
        raise IngestionError(f"Could not open the PDF: {exc}") from exc
    if document.needs_pass:
        document.close()
        raise IngestionError("The PDF is password-protected; upload an unlocked copy.")
    if document.page_count < 1:
        document.close()
        raise IngestionError("The PDF has no pages.")

    result = DocumentAnalysis(file_name=file_name, page_count=document.page_count)
    next_path_id = 0
    ocr_available = use_ocr and bool(shutil.which("tesseract"))
    for page_number, page in enumerate(document, start=1):
        result.page_sizes[page_number] = (float(page.rect.width), float(page.rect.height))
        spans = _span_records(page)
        page_tags: list[TagOccurrence] = []
        for span in spans:
            for tag in find_tag_strings(span["text"]):
                page_tags.append(
                    TagOccurrence(
                        original=tag,
                        normalized=normalize_tag(tag),
                        page=page_number,
                        bbox=span["bbox"],
                        source="selectable text",
                        context=_context_for_bbox(span["bbox"], spans),
                    )
                )

        if use_ocr and (not page_tags or len(" ".join(s["text"] for s in spans)) < 24):
            if ocr_available:
                ocr_tags = _ocr_page(page, page_number)
                for item in ocr_tags:
                    same_text_nearby = any(
                        existing.normalized == item.normalized
                        and math.dist(
                            (
                                (existing.bbox[0] + existing.bbox[2]) / 2,
                                (existing.bbox[1] + existing.bbox[3]) / 2,
                            ),
                            (
                                (item.bbox[0] + item.bbox[2]) / 2,
                                (item.bbox[1] + item.bbox[3]) / 2,
                            ),
                        ) < 4
                        for existing in page_tags
                    )
                    if not same_text_nearby:
                        page_tags.append(item)
            elif not page_tags:
                result.warnings.append(
                    f"Page {page_number}: no selectable tag text was found and Tesseract OCR is unavailable."
                )

        result.tags.extend(page_tags)
        page_segments, next_path_id = _extract_page_vectors(
            page, page_number, next_path_id
        )
        result.segments.extend(page_segments)

        rects = _rectangles(page, page_number)
        for tag in page_tags:
            kind, evidence, confidence = _equipment_kind(tag.original, tag.context)
            if kind:
                result.equipment.append(
                    Equipment(
                        tag=tag.original,
                        normalized=tag.normalized,
                        kind=kind,
                        page=tag.page,
                        bbox=tag.bbox,
                        evidence=evidence,
                        confidence=confidence,
                        context=tag.context,
                        symbol_bbox=_symbol_bounds(tag.bbox, rects),
                    )
                )

    document.close()

    # Collapse repeated detections of the same tag at the same location, but
    # retain distinct labels so duplicates remain visible in the tag report.
    unique_tags: list[TagOccurrence] = []
    for tag in result.tags:
        duplicate = any(
            old.normalized == tag.normalized
            and old.page == tag.page
            and math.dist(
                (
                    (old.bbox[0] + old.bbox[2]) / 2,
                    (old.bbox[1] + old.bbox[3]) / 2,
                ),
                (
                    (tag.bbox[0] + tag.bbox[2]) / 2,
                    (tag.bbox[1] + tag.bbox[3]) / 2,
                ),
            ) < 2.5
            for old in unique_tags
        )
        if not duplicate:
            unique_tags.append(tag)
    result.tags = unique_tags

    unique_equipment: dict[tuple[str, int], Equipment] = {}
    for asset in result.equipment:
        key = (asset.normalized, asset.page)
        prior = unique_equipment.get(key)
        if prior is None or (asset.confidence == "label-confirmed" and prior.confidence != "label-confirmed"):
            unique_equipment[key] = asset
    result.equipment = list(unique_equipment.values())
    if not result.segments:
        result.warnings.append("No vector line segments were extracted; the PDF may be raster-only.")
    return result
